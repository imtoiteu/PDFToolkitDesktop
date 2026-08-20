"""Dựng tệp Word của báo cáo từ mô tả nội dung dạng danh sách."""

from __future__ import annotations

import os
import sys

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docxsvg import add_svg_picture, svg_size  # noqa: E402

FIGDIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "svg"
)
SHOTDIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "screenshots"
)

TEXT_WIDTH_CM = 16.0     # 21 - 3.5 - 1.5


# --------------------------------------------------------------------------
def _set_style(style, *, font="Times New Roman", size=14, bold=False,
               italic=False, align=None, first_line=None, before=0, after=4,
               line=1.2, page_break=False, keep_next=False):
    f = style.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), font)
    pf = style.paragraph_format
    if align is not None:
        pf.alignment = align
    if first_line is not None:
        pf.first_line_indent = first_line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    ppr = style.element.get_or_add_pPr()
    if page_break and ppr.find(qn("w:pageBreakBefore")) is None:
        ppr.append(OxmlElement("w:pageBreakBefore"))
    wc = ppr.find(qn("w:widowControl"))
    if wc is None:
        wc = OxmlElement("w:widowControl")
        ppr.append(wc)
    wc.set(qn("w:val"), "0")


class Report:
    def __init__(self):
        self.doc = docx.Document()
        self._setup()
        self.fig_count = 0
        self.tab_count = 0
        self.placeholders = []
        self.used_figs = []

    # ---------------------------------------------------------------- setup
    def _setup(self):
        d = self.doc
        s = d.sections[0]
        s.page_width = Cm(21)
        s.page_height = Cm(29.7)
        s.left_margin = Cm(3.5)
        s.right_margin = Cm(1.5)
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.0)

        _set_style(d.styles["Normal"], size=14,
                   align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                   first_line=Cm(1.25), after=4, line=1.2)
        _set_style(d.styles["Heading 1"], size=14, bold=True,
                   font="Times New Roman",
                   align=WD_ALIGN_PARAGRAPH.CENTER, first_line=Cm(0),
                   before=2, after=8, line=1.2, page_break=True,
                   keep_next=True)
        _set_style(d.styles["Heading 2"], size=14, bold=True,
                   align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=Cm(1.25),
                   before=6, after=4, line=1.2, keep_next=True)
        _set_style(d.styles["Heading 3"], size=14, bold=True, italic=True,
                   align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=Cm(1.25),
                   before=6, after=4, line=1.2, keep_next=True)
        for h in ("Heading 1", "Heading 2", "Heading 3"):
            d.styles[h].font.color.rgb = RGBColor(0, 0, 0)

        # kiểu riêng cho chú thích hình/bảng
        cap = d.styles.add_style("ChuThich", 1)
        cap.base_style = d.styles["Normal"]
        _set_style(cap, size=13, bold=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER, first_line=Cm(0),
                   before=2, after=8, line=1.0)

        fig = d.styles.add_style("HinhAnh", 1)
        fig.base_style = d.styles["Normal"]
        _set_style(fig, size=14, align=WD_ALIGN_PARAGRAPH.CENTER,
                   first_line=Cm(0), before=6, after=2, line=1.0,
                   keep_next=True)

        bul = d.styles.add_style("GachDau", 1)
        bul.base_style = d.styles["Normal"]
        _set_style(bul, size=14, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                   first_line=Cm(0), before=0, after=3, line=1.2)
        bul.paragraph_format.left_indent = Cm(1.25)

        cell = d.styles.add_style("OBang", 1)
        cell.base_style = d.styles["Normal"]
        _set_style(cell, size=12, align=WD_ALIGN_PARAGRAPH.LEFT,
                   first_line=Cm(0), before=2, after=2, line=1.0)

        note = d.styles.add_style("GhiChu", 1)
        note.base_style = d.styles["Normal"]
        _set_style(note, size=13, italic=True,
                   align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=Cm(1.25),
                   before=2, after=6, line=1.15)

        code = d.styles.add_style("MaNguon", 1)
        code.base_style = d.styles["Normal"]
        _set_style(code, font="Consolas", size=11,
                   align=WD_ALIGN_PARAGRAPH.LEFT, first_line=Cm(0),
                   before=3, after=6, line=1.05)
        code.paragraph_format.left_indent = Cm(1.0)

        ph = d.styles.add_style("ChoAnh", 1)
        ph.base_style = d.styles["Normal"]
        _set_style(ph, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                   first_line=Cm(0), before=4, after=2, line=1.15)

    # ------------------------------------------------------------ primitives
    def h1(self, text, page_break=True):
        p = self.doc.add_paragraph(text, style="Heading 1")
        if not page_break:
            ppr = p._p.get_or_add_pPr()
            el = ppr.find(qn("w:pageBreakBefore"))
            if el is None:
                el = OxmlElement("w:pageBreakBefore")
                ppr.append(el)
            el.set(qn("w:val"), "0")
        return p

    def h2(self, text):
        return self.doc.add_paragraph(text, style="Heading 2")

    def h3(self, text):
        return self.doc.add_paragraph(text, style="Heading 3")

    def p(self, text):
        return self.doc.add_paragraph(text, style="Normal")

    def note(self, text):
        return self.doc.add_paragraph(text, style="GhiChu")

    def code(self, text):
        for line in text.split("\n"):
            self.doc.add_paragraph(line, style="MaNguon")

    def bullets(self, items):
        for it in items:
            self.doc.add_paragraph("–  " + it, style="GachDau")

    def numbered(self, items):
        for i, it in enumerate(items):
            self.doc.add_paragraph("%d)  %s" % (i + 1, it), style="GachDau")

    # ---------------------------------------------------------------- figure
    def figure(self, name, caption, width_cm=None, chapter=None):
        path = os.path.join(FIGDIR, name + ".svg")
        if not os.path.exists(path):
            raise FileNotFoundError(path)
        w_px, h_px = svg_size(path)
        if width_cm is None:
            width_cm = TEXT_WIDTH_CM
            # giữ chiều cao không vượt quá 20 cm
            if h_px / w_px * width_cm > 19.5:
                width_cm = 19.5 * w_px / h_px
        par = self.doc.add_paragraph(style="HinhAnh")
        add_svg_picture(self.doc, par, path, Cm(width_cm))
        self.doc.add_paragraph(caption, style="ChuThich")
        self.used_figs.append((name, caption))
        return par

    def screenshot(self, filename, caption, width_cm=None):
        path = os.path.join(SHOTDIR, filename)
        par = self.doc.add_paragraph(style="HinhAnh")
        run = par.add_run()
        run.add_picture(path, width=Cm(width_cm or TEXT_WIDTH_CM))
        self.doc.add_paragraph(caption, style="ChuThich")
        return par

    def placeholder(self, fig_label, caption, what, steps):
        """Khung giữ chỗ cho ảnh chụp màn hình cần bổ sung thủ công."""
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        cell.width = Cm(TEXT_WIDTH_CM)
        _shade(cell, "FFF7E6")
        _borders(tbl, color="D97706", sz=12)
        c0 = cell.paragraphs[0]
        c0.style = self.doc.styles["ChoAnh"]
        c0.add_run("[CHỖ DÀNH CHO ẢNH CHỤP MÀN HÌNH – %s]" % fig_label)
        p2 = cell.add_paragraph(style="ChoAnh")
        p2.add_run(what).italic = True
        p3 = cell.add_paragraph(style="OBang")
        p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p3.add_run("Thao tác cần thực hiện để chụp:").bold = True
        for i, s in enumerate(steps):
            ps = cell.add_paragraph(style="OBang")
            ps.paragraph_format.left_indent = Cm(0.6)
            ps.add_run("%d. %s" % (i + 1, s))
        self.doc.add_paragraph(caption, style="ChuThich")
        self.placeholders.append((fig_label, caption, what, steps))

    # ----------------------------------------------------------------- table
    def table(self, caption, headers, rows, widths=None, font_size=12):
        self.doc.add_paragraph(caption, style="ChuThich")
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _borders(t)
        hdr = t.rows[0]
        for i, h in enumerate(headers):
            c = hdr.cells[i]
            _shade(c, "E8EEF7")
            par = c.paragraphs[0]
            par.style = self.doc.styles["OBang"]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(h)
            r.bold = True
            r.font.size = Pt(font_size)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                par = cells[i].paragraphs[0]
                par.style = self.doc.styles["OBang"]
                r = par.add_run(str(v))
                r.font.size = Pt(font_size)
        if widths:
            total = sum(widths)
            for r_ in t.rows:
                for i, c in enumerate(r_.cells):
                    c.width = Cm(TEXT_WIDTH_CM * widths[i] / total)
        _repeat_header(t)
        self.doc.add_paragraph("", style="OBang")
        return t

    # ------------------------------------------------------------------ misc
    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def toc_field(self):
        par = self.doc.add_paragraph(style="Normal")
        par.paragraph_format.first_line_indent = Cm(0)
        fld = parse_xml(
            '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            "<w:r><w:t>Nhấn chuột phải và chọn “Update Field” để sinh mục lục.</w:t></w:r>"
            '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
        )
        par._p.addnext(fld)

    def page_numbers(self):
        W = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
             '2006/main"')
        sec = self.doc.sections[0]
        par = sec.footer.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.style = self.doc.styles["Normal"]
        par.paragraph_format.first_line_indent = Cm(0)
        run = par.add_run()
        run._r.append(parse_xml('<w:fldChar %s w:fldCharType="begin"/>' % W))
        run._r.append(parse_xml(
            '<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % W))
        run._r.append(parse_xml('<w:fldChar %s w:fldCharType="end"/>' % W))

    def save(self, path):
        self.doc.save(path)


# --------------------------------------------------------------------------
def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _borders(table, color="7A7A7A", sz=6):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def _repeat_header(table):
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)
